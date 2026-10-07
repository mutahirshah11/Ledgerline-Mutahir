"""F-02 input validation: pure contract checks over file bytes (Section 7, Section 14).

Reads only; never changes data (F-02.R5). Every problem from every file is returned, in source
order REGISTER, ERP, SHEET (`order.SOURCE_ORDER`) and file order within a source (F-02.R4).
"""

import csv
import io
import re
import struct
import zipfile
from collections.abc import Sequence
from datetime import date
from decimal import Decimal, InvalidOperation
from typing import NamedTuple

import openpyxl

from app.domain.enums import ProblemCode, Source
from app.domain.models import DomainModel

ERP_COLUMNS: tuple[str, ...] = (
    "grn_no",
    "grn_date",
    "supplier_code",
    "supplier_name",
    "item_desc",
    "lot_no",
    "qty",
    "uom",
    "entered_by",
    "entered_at",
)
"""The ten required ERP columns, exact names (Section 7)."""

SHEET_HEADERS: tuple[str, ...] = (
    "Date",
    "GRN",
    "Supplier",
    "Material",
    "Lot",
    "Qty Received",
    "UOM",
    "Remarks",
)
"""The eight required store sheet headers in row 1, exact text (Section 7)."""

MIN_IMAGE_PX = 800
MAX_PAGES = 10

_MONTHS = ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")
_ISO_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_PNG_SIGNATURE = b"\x89PNG\r\n\x1a\n"
# JPEG start-of-frame markers carry the frame size; C4 (DHT), C8 (JPG), CC (DAC) are not frames.
_JPEG_SOF = frozenset(range(0xC0, 0xD0)) - {0xC4, 0xC8, 0xCC}
_JPEG_STANDALONE = frozenset({0x01, *range(0xD0, 0xD8)})


class ProblemLocation(DomainModel):
    """Where a problem is: ERP data `line` (1 = first data row) and `column`; sheet `row`,
    `column` (header) and `worksheet`; register `page_no` (data-model.md, Validation)."""

    line: int | None = None
    row: int | None = None
    column: str | None = None
    worksheet: str | None = None
    page_no: int | None = None


class ValidationProblem(DomainModel):
    """One broken contract rule, in plain language for SCR-03 (F-02.R4)."""

    source: Source
    code: ProblemCode
    message: str
    location: ProblemLocation | None = None


class RegisterFile(NamedTuple):
    """One uploaded register page: its page number, original file name and bytes."""

    page_no: int
    filename: str
    data: bytes


def _problem(
    source: Source, code: ProblemCode, message: str, **location: int | str
) -> ValidationProblem:
    loc = ProblemLocation.model_validate(location) if location else None
    return ValidationProblem(source=source, code=code, message=message, location=loc)


# --- R1 ERP -------------------------------------------------------------------------------


def _is_decimal(value: str) -> bool:
    try:
        return Decimal(value).is_finite()
    except InvalidOperation:
        return False


def _is_iso_date(value: str) -> bool:
    if not _ISO_DATE.match(value):
        return False
    try:
        date.fromisoformat(value)
    except ValueError:
        return False
    return True


def check_erp(data: bytes) -> list[ValidationProblem]:
    """F-02.R1: UTF-8 (BOM allowed), quoted CSV, the ten columns, ≥ 1 data row, decimal `qty`
    and ISO `grn_date` on every data row (line 1 = first data row)."""
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        msg = f"ERP: file is not UTF-8 text (bad byte at position {exc.start})"
        return [_problem(Source.ERP, ProblemCode.ENCODING, msg)]
    reader = csv.reader(io.StringIO(text, newline=""))
    header = next(reader, [])
    problems = [
        _problem(Source.ERP, ProblemCode.MISSING_COLUMN, f"ERP: column {c} is missing", column=c)
        for c in ERP_COLUMNS
        if c not in header
    ]
    rows = [r for r in reader if any(v.strip() for v in r)]
    if not rows:
        problems.append(_problem(Source.ERP, ProblemCode.NO_DATA_ROWS, "ERP: no data rows"))
    checks = (
        ("qty", ProblemCode.BAD_NUMBER, _is_decimal, "is not a number"),
        ("grn_date", ProblemCode.BAD_DATE, _is_iso_date, "is not a date YYYY-MM-DD"),
    )
    index = {name: header.index(name) for name, *_ in checks if name in header}
    for line, row in enumerate(rows, start=1):
        for name, code, ok, what in checks:
            if name not in index:
                continue
            value = row[index[name]] if index[name] < len(row) else ""
            if not ok(value):
                msg = f'ERP line {line}: {name} "{value}" {what}'
                problems.append(_problem(Source.ERP, code, msg, line=line, column=name))
    return problems


# --- R2 sheet -----------------------------------------------------------------------------


def sheet_name(month: str) -> str:
    """The worksheet name for run month `YYYY-MM`: `2026-09` → `Sep-2026` (Section 7)."""
    year, mon = month.split("-")
    return f"{_MONTHS[int(mon) - 1]}-{year}"


def check_sheet(month: str, data: bytes) -> list[ValidationProblem]:
    """F-02.R2: openpyxl read-only with values; worksheet `Mon-YYYY` for the run month; the
    eight headers in row 1; ≥ 1 data row. Cell values are not parsed (F-02.AC6: F-04's job)."""
    name = sheet_name(month)
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (zipfile.BadZipFile, OSError, KeyError, ValueError):
        msg = "Store sheet: file is not a readable .xlsx workbook"
        return [_problem(Source.SHEET, ProblemCode.ENCODING, msg)]
    try:
        if name not in wb.sheetnames:
            found = ", ".join(wb.sheetnames) or "none"
            msg = f"Store sheet: no worksheet named {name} (found: {found})"
            return [_problem(Source.SHEET, ProblemCode.MISSING_WORKSHEET, msg, worksheet=name)]
        rows = wb[name].iter_rows(values_only=True)
        header = next(rows, ())
        problems = [
            _problem(
                Source.SHEET,
                ProblemCode.MISSING_HEADER,
                f"Store sheet: header {h} is missing from row 1",
                row=1,
                column=h,
                worksheet=name,
            )
            for h in SHEET_HEADERS
            if h not in header
        ]
        if not any(any(v is not None and str(v).strip() for v in r) for r in rows):
            msg = f"Store sheet: worksheet {name} has no data rows"
            problems.append(_problem(Source.SHEET, ProblemCode.NO_DATA_ROWS, msg, worksheet=name))
        return problems
    finally:
        wb.close()


# --- R3 register --------------------------------------------------------------------------


def _png_size(data: bytes) -> tuple[int, int] | None:
    if len(data) < 24 or not data.startswith(_PNG_SIGNATURE) or data[12:16] != b"IHDR":
        return None
    width, height = struct.unpack(">II", data[16:24])
    return width, height


def _jpeg_size(data: bytes) -> tuple[int, int] | None:
    if not data.startswith(b"\xff\xd8"):
        return None
    i = 2
    while i + 4 <= len(data):
        if data[i] != 0xFF:
            return None
        marker = data[i + 1]
        if marker == 0xFF:  # fill byte
            i += 1
            continue
        if marker in _JPEG_STANDALONE:
            i += 2
            continue
        if marker in (0xD9, 0xDA):  # end of image, or scan data before any frame
            return None
        (length,) = struct.unpack(">H", data[i + 2 : i + 4])
        if marker in _JPEG_SOF:
            if i + 9 > len(data):
                return None
            height, width = struct.unpack(">HH", data[i + 5 : i + 9])
            return width, height
        i += 2 + length
    return None


def image_size(data: bytes) -> tuple[int, int] | None:
    """Width and height from a PNG IHDR or JPEG SOF header, standard library only (research
    R6); None when the bytes are neither or the header can't be parsed."""
    size = _png_size(data) or _jpeg_size(data)
    if size is None or 0 in size:
        return None
    return size


def check_register(files: Sequence[RegisterFile]) -> list[ValidationProblem]:
    """F-02.R3: at most ten pages; each opens as a PNG or JPEG at least 800 × 800 px.
    Problems follow the order of `files` (callers pass them by page number)."""
    problems: list[ValidationProblem] = []
    if len(files) > MAX_PAGES:
        msg = f"Register: {len(files)} pages uploaded; at most {MAX_PAGES} allowed"
        problems.append(_problem(Source.REGISTER, ProblemCode.TOO_MANY_PAGES, msg))
    for f in files:
        where = f"Register page {f.page_no} ({f.filename})"
        size = image_size(f.data)
        if size is None:
            msg = f"{where}: not a readable JPEG or PNG image"
            problems.append(
                _problem(Source.REGISTER, ProblemCode.IMAGE_UNREADABLE, msg, page_no=f.page_no)
            )
        elif min(size) < MIN_IMAGE_PX:
            msg = (
                f"{where}: image is {size[0]} x {size[1]} px; "
                f"at least {MIN_IMAGE_PX} x {MIN_IMAGE_PX} needed"
            )
            problems.append(
                _problem(Source.REGISTER, ProblemCode.IMAGE_TOO_SMALL, msg, page_no=f.page_no)
            )
    return problems


# --- R4 all files -------------------------------------------------------------------------


def validate_files(
    month: str, erp_bytes: bytes, sheet_bytes: bytes, register_files: Sequence[RegisterFile]
) -> list[ValidationProblem]:
    """Every problem across all files, REGISTER then ERP then SHEET (F-02.R4); [] = pass."""
    return [
        *check_register(register_files),
        *check_erp(erp_bytes),
        *check_sheet(month, sheet_bytes),
    ]
